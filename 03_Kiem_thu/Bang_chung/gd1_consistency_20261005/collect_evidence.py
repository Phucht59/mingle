"""Read-only extraction for the 2026-10-05 GD1 audit. No workbook authoring."""
from pathlib import Path
import csv
import hashlib
import json
import re
import subprocess
from datetime import datetime, timezone
import openpyxl

R = Path('C:/Mingo')
E = R / '03_Kiem_thu/Bang_chung/gd1_consistency_20261005'
U = R / '02_Tai_lieu_du_an/04_Thiet_ke_san_pham/ux_ui/phase2/rebaseline_v2'
def save(name, value):
    (E/name).write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()
def csvread(name):
    return list(csv.DictReader((U/name).open(encoding='utf-8-sig', newline='')))

inventory = []
for root in ['01_San_pham','02_Tai_lieu_du_an','04_Van_hanh','.github']:
    for p in (R/root).rglob('*'):
        if not p.is_file() or any(x in p.parts for x in ['build','node_modules','.dart_tool','ephemeral','__pycache__','.pytest_cache','.ruff_cache']):
            continue
        if p.suffix.lower() not in ['.md','.json','.csv','.yaml','.sql','.py','.dart','.xlsx','.toml']:
            continue
        inventory.append({'path':p.relative_to(R).as_posix(),'sha256':sha(p),'bytes':p.stat().st_size})
save('active_source_inventory_before.json',inventory)
save('git_identity.json',{
    'local_date':'2026-10-05','timezone':'Asia/Saigon','captured_utc':datetime.now(timezone.utc).isoformat(),
    'branch':subprocess.check_output(['git','branch','--show-current'],cwd=R,text=True).strip(),
    'commit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=R,text=True).strip(),
    'remote':subprocess.check_output(['git','remote','get-url','origin'],cwd=R,text=True).strip(),
    'source_boundary':'dirty working tree; NOT a clean committed candidate',
})

p=R/'02_Tai_lieu_du_an/07_Tien_do_du_an/Quan_ly_du_an_Mingo.xlsx'
w=openpyxl.load_workbook(p,read_only=True,data_only=False)
d=openpyxl.load_workbook(p,read_only=True,data_only=True)
workbook={'path':p.relative_to(R).as_posix(),'sha256':sha(p),'sheets':[]}
for s in w:
    rows=[]
    for row in s.iter_rows():
        cells=[{'cell':c.coordinate,'value':c.value,'cached':d[s.title][c.coordinate].value if c.data_type=='f' else None} for c in row if c.value is not None]
        if cells: rows.append(cells)
    workbook['sheets'].append({'sheet':s.title,'nonempty_rows':rows})
save('available_management_workbook_extract.json',workbook)
w.close();d.close()

catalog=json.loads((U/'SCREEN_CATALOG.json').read_text(encoding='utf-8-sig'))
qa=json.loads((R/'03_Kiem_thu/QA_QC/phase2/qa_cases.json').read_text(encoding='utf-8-sig'))
prd=csvread('PRD_TRACEABILITY_V2.csv'); trace=csvread('SCREEN_STATE_TRACEABILITY.csv'); scopes=csvread('SCREEN_SCOPE_REGISTER.csv')
screen_ids={x['id'] for x in catalog}; qa_ids={x['id'] for x in qa}
expected={(s['id'],state) for s in catalog for state in s['states']}
seen=[(x['screen_id'],x['state']) for x in trace]
gaps=[]
for n,x in enumerate(prd,2):
    for sid in x['screens'].split(','):
        if sid not in screen_ids:gaps.append({'file':'PRD_TRACEABILITY_V2.csv','row':n,'ref':sid,'kind':'screen'})
    for qid in x['qa'].split(','):
        if qid not in qa_ids:gaps.append({'file':'PRD_TRACEABILITY_V2.csv','row':n,'ref':qid,'kind':'QA'})
for n,x in enumerate(trace,2):
    for qid in re.findall(r'(?:QA-[A-Z]+|QC)-\d+',x['inherited_qa_case_ids']):
        if qid not in qa_ids:gaps.append({'file':'SCREEN_STATE_TRACEABILITY.csv','row':n,'ref':qid,'kind':'QA'})
    if not (R/x['screenshot']).is_file():gaps.append({'file':'SCREEN_STATE_TRACEABILITY.csv','row':n,'ref':x['screenshot'],'kind':'screenshot'})
save('structural_audit.json',{
    'catalog_screens':len(screen_ids),'catalog_states':len(expected),'trace_rows':len(trace),
    'missing_states':sorted(expected-set(seen)),'extra_states':sorted(set(seen)-expected),'duplicate_state_rows':len(seen)-len(set(seen)),
    'requirements':len(prd),'duplicate_prd_ids':len(prd)-len({x['prd'] for x in prd}),
    'dangling_refs':gaps,'scope_rows':len(scopes),'missing_scope':sorted(screen_ids-{x['screen_id'] for x in scopes}),
    'future_routes':[x['screen_id'] for x in scopes if x['category']=='FUTURE'],
    'screens_without_direct_PRD_mapping':sorted(screen_ids-{s for x in prd for s in x['screens'].split(',')}),
    'note':'Structural existence only; no assertion of semantic completeness or current runtime QA PASS.',
    'qa_priority_by_requirement':{x['prd']:[{'id':q['id'],'priority':q['priority'],'expected':q['expected'],'title':q['title']} for q in qa if q['id'] in x['qa'].split(',')] for x in prd},
})
save('prd_trace_source.json',prd)

f=E/'flutter_selected.jsonl'
if f.exists():
    events=[]
    for line in f.read_text(encoding='utf-8-sig',errors='replace').splitlines():
        try:events.append(json.loads(line))
        except json.JSONDecodeError:pass
    starts={e['test']['id']:e['test']['name'] for e in events if e.get('type')=='testStart'}
    done=[e for e in events if e.get('type')=='testDone' and not e.get('hidden')]
    save('flutter_selected_summary.json',{'passed':sum(e.get('result')=='success' and not e.get('skipped') for e in done),'failed':sum(e.get('result') in ['failure','error'] for e in done),'skipped':sum(bool(e.get('skipped')) for e in done),'suite_done':[e for e in events if e.get('type')=='done'],'tests':[{'name':starts.get(e['testID']),'result':e.get('result'),'skipped':e.get('skipped')} for e in done]})
print(json.dumps({'inventory_files':len(inventory),'management_workbook_sheets':len(workbook['sheets']),'screens':len(screen_ids),'states':len(expected),'dangling_refs':gaps,'future_routes':[x['screen_id'] for x in scopes if x['category']=='FUTURE']},ensure_ascii=False))
