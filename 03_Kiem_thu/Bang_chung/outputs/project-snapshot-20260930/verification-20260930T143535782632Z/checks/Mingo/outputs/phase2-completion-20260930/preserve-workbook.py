from pathlib import Path
import json, re, hashlib, zipfile
from xml.sax.saxutils import escape
import openpyxl

folder=Path(r'C:\Mingo\outputs\phase2-completion-20260930')
source=Path(r'C:\Mingo\06_quality\phase2\final_gate\2026-09-30\inputs\Mingo_Phase2_R1_Independent_QA_Control_2026-09-29.xlsx')
output=folder/'Mingo_Phase2_QA_Control_Dashboard_Corrected_2026-09-30.xlsx'
edits=json.loads((folder/'dashboard-edits.json').read_text(encoding='utf-8'))
original_hash=hashlib.sha256(source.read_bytes()).hexdigest()
assert original_hash=='6c3792d76628e894c46e814db933714a63b6669fd6bd936f1ac3b4d575438222'
target='xl/worksheets/sheet1.xml'
with zipfile.ZipFile(source) as z:
    original={n:z.read(n) for n in z.namelist()}
    xml=original[target].decode('utf-8')
    for cell,item in edits.items():
        pattern=rf'<(?P<tag>(?:[A-Za-z_][\w.-]*:)?c)\b(?P<attrs>[^>]*\br="{re.escape(cell)}"[^>]*)>(?P<body>.*?)</(?P=tag)>'
        matches=list(re.finditer(pattern,xml,re.S))
        assert len(matches)==1,(cell,len(matches))
        tag=matches[0].group('tag')
        prefix=tag[:-1]
        attrs=re.sub(r'\s+t="[^"]*"','',matches[0].group('attrs'))
        value=item['value']
        if item['formula']:
            body=f'<{prefix}f>'+escape(item['formula'].removeprefix('='))+f'</{prefix}f><{prefix}v>'+str(value)+f'</{prefix}v>'
        elif isinstance(value,(int,float)):
            body=f'<{prefix}v>'+str(value)+f'</{prefix}v>'
        else:
            attrs+=' t="inlineStr"'
            body=f'<{prefix}is><{prefix}t>'+escape(str(value))+f'</{prefix}t></{prefix}is>'
        replacement='<'+tag+attrs+'>'+body+'</'+tag+'>'
        xml=xml[:matches[0].start()]+replacement+xml[matches[0].end():]
    with zipfile.ZipFile(output,'w') as dst:
        for info in z.infolist():dst.writestr(info,xml.encode('utf-8') if info.filename==target else original[info.filename])
with zipfile.ZipFile(output) as z:
    changed_parts=[n for n in original if z.read(n)!=original[n]]
assert changed_parts==[target],changed_parts
before=openpyxl.load_workbook(source,data_only=False)
after=openpyxl.load_workbook(output,data_only=False)
cached=openpyxl.load_workbook(output,data_only=True)
assert before.sheetnames==after.sheetnames
changed_cells=[]
for name in before.sheetnames:
    old,new=before[name],after[name]
    assert old.max_row==new.max_row and old.max_column==new.max_column
    assert str(old.merged_cells)==str(new.merged_cells)
    assert old.freeze_panes==new.freeze_panes
    assert str(old.data_validations)==str(new.data_validations)
    for row in old:
        for cell in row:
            dest=new[cell.coordinate]
            assert cell._style==dest._style,(name,cell.coordinate,'style')
            if cell.value!=dest.value:
                assert name=='00_Dashboard' and cell.coordinate in edits,(name,cell.coordinate)
                changed_cells.append(cell.coordinate)
            assert cell.comment==dest.comment,(name,cell.coordinate,'comment')
assert set(changed_cells)==set(edits)
qa=[r for r in after['04_QA_Test_Cases'].iter_rows(min_row=5,values_only=True) if r[0]]
totals={'unique_cases':len(set(r[0] for r in qa)),'p0':sum(r[3]=='P0' for r in qa),'pass':sum(r[10]=='PASS' for r in qa)}
assert totals=={'unique_cases':94,'p0':44,'pass':94},totals
for cell,expected in {'E5':94,'E6':44,'E7':94,'E8':0,'B10':94}.items():assert cached['00_Dashboard'][cell].value==expected
assert cached['00_Dashboard']['B6'].value=='HOLD / NOT PASSED'
report={'source':str(source),'output':str(output),'source_sha256':original_hash,'output_sha256':hashlib.sha256(output.read_bytes()).hexdigest(),'changed_parts':changed_parts,'changed_cells':changed_cells,'totals':totals,'unchanged_sheets':13,'all_non_dashboard_zip_parts_byte_identical':True,'all_sheet_cell_styles_preserved':True,'signoff_sheet_byte_identical':True,'gate':'HOLD / NOT PASSED','independent_report_confirmation':'PENDING','visual_check':'native Excel required; artifact renderer exceeded two bounded attempts'}
(folder/'workbook-correction-verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
(folder/(output.name+'.sha256')).write_text(report['output_sha256']+'  '+output.name+'\n',encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
