import openpyxl,json,warnings,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
warnings.filterwarnings('ignore')
d=Path('C:/Mingo/outputs/mingo-ui-20260927')
original=openpyxl.load_workbook('C:/Mingo/Manage Project.xlsx')
out=openpyxl.load_workbook(d/'Manage Project - Redesigned.xlsx')
repaired=openpyxl.load_workbook(d/'native-repaired.xlsx')
for a,b in zip(out,repaired):
 changes=[]
 for row in a:
  for c in row:
   if c.value!=b[c.coordinate].value: changes.append((c.coordinate,str(c.value),str(b[c.coordinate].value)))
 print(a.title,'repair changes',changes[:8], 'tables',list(a.tables),list(b.tables),'dv',len(a.data_validations.dataValidation),len(b.data_validations.dataValidation),'cf',len(a.conditional_formatting),len(b.conditional_formatting))
for a,b in zip(list(original)[2:],list(out)[2:]):
 changes=[]
 for row in a:
  for c in row:
   if c.value!=b[c.coordinate].value:changes.append((c.coordinate,str(c.value),str(b[c.coordinate].value)))
 print(a.title,'source changes',changes[:12])
z1=zipfile.ZipFile(d/'Manage Project - Redesigned.xlsx');z2=zipfile.ZipFile(d/'native-repaired.xlsx')
ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
for n in ['xl/styles.xml']+[f'xl/worksheets/sheet{i}.xml' for i in range(1,8)]:
 for z in [z1,z2]:
  t=ET.fromstring(z.read(n));print(n,'original' if z==z1 else 'repaired',[(e.tag.split('}')[-1],len(e),e.attrib) for e in t if e.tag.split('}')[-1] not in ['sheetData','cols','mergeCells','conditionalFormatting']])
