from pathlib import Path
import openpyxl,json,datetime,warnings
import pypdfium2 as pdf
from PIL import Image,ImageOps,ImageDraw
warnings.filterwarnings('ignore')
root=Path('C:/Mingo/outputs/mingo-ui-v2')
m=json.loads((root/'manifest.json').read_text(encoding='utf8'))
book=openpyxl.load_workbook(root/'Mingo - Workspace.xlsx')
vals=openpyxl.load_workbook(root/'Mingo - Workspace.xlsx',data_only=True)
def expected(v):
 if isinstance(v,str) and v.startswith('='):
  return v.replace("'04_CONG_VIEC'!$A$5:$F$200,6,FALSE","'02_Cong_viec'!$B$9:$F$204,5,FALSE").replace("'04_CONG_VIEC'!$B$5:$B$200","'02_Cong_viec'!$C$9:$C$204").replace("'04_CONG_VIEC'!$F$5:$F$200","'02_Cong_viec'!$F$9:$F$204")
 if isinstance(v,str) and len(v)==19 and v.endswith(' 00:00:00'):return datetime.datetime.fromisoformat(v)
 if isinstance(v,str):return v.replace('04_CONG_VIEC','02_Cong_viec').replace('02_SAN_PHAM','08_San_pham')
 return v
bad=[]
for a in m['audit']:
 if 'value' not in a:continue
 actual=book[a['target']][a['at']].value
 if actual!=expected(a['value']):bad.append((a['source'],a['cell'],a['target'],a['at'],actual,a['value']))
assert not bad,bad
errors=[(s.title,c.coordinate,c.value) for s in vals for row in s for c in row if c.data_type=='e']
assert not errors,errors
assert vals['00_Tong_quan']['E8'].value==0
assert vals['00_Tong_quan']['H8'].value==8
assert book['02_Cong_viec']['F9'].value=='ĐANG LÀM'
assert len(book.sheetnames)==22
assert all(t.autoFilter is not None for s in book for t in s.tables.values())
print('PASS:',len(m['audit']),'mapped source fields; 22 sheets; no formula errors; live values and filters verified.')
thumbs=[]
for sm in sorted(m['sheets'],key=lambda x:x['name']):
 p=root/(sm['name']+'-final.pdf')
 if not p.exists():p=root/(sm['name']+'.pdf')
 with pdf.PdfDocument(p) as doc:
  page=doc[0];im=page.render(scale=1.4).to_pil();page.close()
  # Crop only surrounding page whitespace for inspection, retaining all worksheet content.
  import PIL.ImageChops as ch
  bbox=ch.difference(im.convert('RGB'),Image.new('RGB',im.size,'white')).getbbox()
  if bbox:im=im.crop((max(0,bbox[0]-15),max(0,bbox[1]-15),min(im.width,bbox[2]+15),min(im.height,bbox[3]+15)))
  im.save(root/(sm['name']+'.png'))
  thumb=Image.new('RGB',(560,390),'white');small=im.copy();small.thumbnail((540,350));thumb.paste(small,((560-small.width)//2,28));ImageDraw.Draw(thumb).text((12,8),sm['name'],fill='black');thumbs.append(thumb)
for start in range(0,len(thumbs),6):
 contact=Image.new('RGB',(1680,780),'#e6e6e6')
 for j,im in enumerate(thumbs[start:start+6]):contact.paste(im,((j%3)*560,(j//3)*390))
 contact.save(root/f'contact-{start//6+1}.png')
